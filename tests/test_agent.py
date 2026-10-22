"""Offline tests: a scripted fake LLM stands in for Groq/Ollama, so these run in CI without keys."""
from __future__ import annotations

from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import prompts
from app.agent import DocumentAgent
from app.llm import LLMError, OpenAICompatibleLLM, _extract_json
from app.main import app, get_agent
from app.schemas import CostEstimate, Interpretation, Plan, Review, Section

COMPLEX_REQUEST = ("We need a plan to move our CRM to the cloud. Budget is tight but leadership wants it "
                   "done in 6 weeks with zero downtime. Not sure how many users we have.")


class FakeLLM:
    """Returns canned, schema-valid objects and records every call."""

    def __init__(self, *, accept=True, fail_section=None, fail_once=None, review_fails_first=True):
        self.calls: list[tuple[str, str]] = []
        self.accept = accept
        self.fail_section = fail_section
        self.fail_once = fail_once
        self.review_fails_first = review_fails_first
        self.reviews = 0

    def generate(self, system, user, schema):
        self.calls.append((schema.__name__, user))
        if schema is Interpretation:
            return Interpretation(
                is_document_request=self.accept, rejection_reason=None if self.accept else "Not a document",
                doc_type="project_plan", title="CRM Cloud Migration Plan", audience="CTO",
                objective="Migrate CRM to cloud", key_facts=["6-week deadline", "zero downtime"],
                missing_info=["number of CRM users"], conflicts=["tight budget vs 6 weeks vs zero downtime"],
                assumptions=["~500 users", "SaaS-compatible data model"])
        if schema is Plan:
            return Plan(reasoning="Standard project plan",
                        key_decisions=["Phase 1 (weeks 1-2): assessment", "Cut-over in week 6"], tasks=[
                {"id": 1, "description": "Summary", "tool": "executive_summary", "section_title": "Executive Summary"},
                {"id": 2, "description": "Scope", "tool": "draft_section", "section_title": "Scope and Approach"},
                {"id": 3, "description": "Timeline", "tool": "build_table", "section_title": "Timeline"},
                {"id": 4, "description": "Risks", "tool": "draft_section", "section_title": "Risks"},
                {"id": 5, "description": "Budget", "tool": "estimate_costs", "section_title": "Budget Estimate"},
            ])
        if schema is Section:
            title = user.split("SECTION TO WRITE: ")[1].split("\n")[0]
            if title == self.fail_section:
                raise RuntimeError("simulated tool failure")
            if title == self.fail_once:
                self.fail_once = None
                raise RuntimeError("simulated transient failure")
            if system == prompts.TABLE_SYSTEM:
                return Section(title=title, blocks=[
                    {"type": "paragraph", "text": "Phased plan (**estimates**):"},
                    {"type": "table", "headers": ["Week", "Milestone", "Owner"],
                     "rows": [[1, "Assessment"], ["3-6", "Cutover", "Infra", "extra"]]}])
            return Section(title=title, blocks=[
                {"type": "paragraph", "text": f"## Content for {title} with **bold** text."},
                {"type": "bullets", "items": ["- point one", "point two"]}])
        if schema is CostEstimate:
            return CostEstimate(line_items=[
                {"role": "project manager", "activity": "Coordination", "days": 10},
                {"role": "Cloud Engineer", "activity": "Migration", "days": 20}],
                other_items=[{"item": "migration tooling licence (per month)", "quantity": 2}],
                notes="Lean team over 6 weeks. Total is $48,480 plus 10% contingency = $50,000.")
        if schema is Review:
            self.reviews += 1
            if self.review_fails_first and self.reviews == 1:
                return Review(passed=False, score=6, issues=[
                    {"section_title": "Risks", "problem": "Conflict not resolved",
                     "fix_instruction": "Add phased trade-off"},
                    {"section_title": "Cost Estimate", "problem": "No budget section",
                     "fix_instruction": "Add an estimated budget"}])
            return Review(passed=True, score=9, issues=[])
        raise AssertionError(f"unexpected schema {schema}")


@pytest.fixture
def make_client(tmp_path):
    def _make(llm):
        app.dependency_overrides[get_agent] = lambda: DocumentAgent(llm, tmp_path, max_revisions=1)
        return TestClient(app)
    yield _make
    app.dependency_overrides.clear()


def test_end_to_end_plan_reflect_and_docx(make_client, tmp_path, monkeypatch):
    llm = FakeLLM()
    client = make_client(llm)
    r = client.post("/agent", json={"request": COMPLEX_REQUEST})
    assert r.status_code == 200, r.text
    body = r.json()

    tools = [t["tool"] for t in body["plan"]]
    # guard appended the assumptions register; reflection inserted the missing section before it
    assert tools[-1] == "assumptions_register"
    assert "Cost Estimate" in [t["section_title"] for t in body["plan"]]
    assert all(t["status"] == "done" for t in body["plan"])

    refl = body["reflection"]
    assert refl["initial_score"] == 6 and refl["final_score"] == 9
    assert {"Risks", "Cost Estimate", "Executive Summary"} <= set(refl["sections_revised"])
    assert body["status"] == "completed"
    assert [s["step"] for s in body["execution_log"]] == ["analyze", "plan", "execute", "reflect", "compile_docx"]

    # content sections ran (concurrently, any order) before the deferred executive summary
    content_calls = [u.split("SECTION TO WRITE: ")[1].split("\n")[0]
                     for n, u in llm.calls if n in ("Section", "CostEstimate")]
    assert set(content_calls[:4]) == {"Scope and Approach", "Timeline", "Risks", "Budget Estimate"}
    assert content_calls[4] == "Executive Summary"

    # download + inspect the .docx
    monkeypatch.setattr("app.main.settings", SimpleNamespace(output_dir=tmp_path))
    d = client.get(body["document_url"])
    assert d.status_code == 200
    doc = Document(tmp_path / f"{body['request_id']}.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "CRM Cloud Migration Plan" in text
    assert "**" not in text and "## " not in text and "- point one" not in text  # markdown stripped
    assert any("~500 users" in c.text for t in doc.tables for row in t.rows for c in row.cells)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Executive Summary", "Scope and Approach", "Timeline", "Risks",
                        "Budget Estimate", "Cost Estimate", "Assumptions and Open Questions"]

    # costs computed in code: labour 10*600 + 20*750 = 21,000; licence 2*1,500 = 3,000;
    # subtotal 24,000 + 15% contingency 3,600 = 27,600
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Project Manager" in cells  # role normalised to rate-card spelling
    assert "Migration tooling licence (per month)" in cells
    assert {"USD 24,000", "USD 3,600", "USD 27,600"} <= set(cells)
    # the model's own arithmetic in its notes must not reach the document
    assert "50,000" not in text and "48,480" not in text and "Lean team over 6 weeks." in text
    assert body["key_decisions"] == ["Phase 1 (weeks 1-2): assessment", "Cut-over in week 6"]

    # usage metering + persisted trace
    assert body["usage"]["llm_calls"] == len(llm.calls)
    run = client.get(f"/runs/{body['request_id']}")
    assert run.status_code == 200 and run.json()["plan"] == body["plan"]


def test_parallel_sections_share_key_decisions_and_outline(tmp_path):
    """Sections are drafted in parallel, so consistency depends on the shared brief."""
    llm = FakeLLM(review_fails_first=False)
    DocumentAgent(llm, tmp_path).run(COMPLEX_REQUEST)
    section_prompts = [u for n, u in llm.calls if n in ("Section", "CostEstimate")]
    review_prompt = next(u for n, u in llm.calls if n == "Review")
    for prompt in section_prompts + [review_prompt]:
        assert "Cut-over in week 6" in prompt
    assert all("Budget Estimate" in u and "Project Manager" in u for u in section_prompts)


def test_multi_point_table_cells_render_one_per_line(tmp_path):
    from app.docx_builder import _clean
    assert _clean("- Encrypt data\n- Rollback snapshot") == "Encrypt data\nRollback snapshot"


def test_reviewer_sees_whole_long_document(tmp_path):
    """Regression: the reviewer input used to be JSON cut at 16k chars, hiding the ending."""
    from app.tools import render_sections_text
    long = [Section(title=f"S{i}", blocks=[{"type": "paragraph", "text": "x" * 3000}]) for i in range(8)]
    long.append(Section(title="Final Section", blocks=[{"type": "paragraph", "text": "END-MARKER"}]))
    text = render_sections_text(long)
    assert len(text) > 16_000 and "END-MARKER" in text and "TRUNCATED" not in text


def test_remaining_issues_recorded_when_recheck_still_fails(tmp_path):
    class StubbornReviewer(FakeLLM):
        def generate(self, system, user, schema):
            self.calls.append((schema.__name__, user))
            return Review(passed=False, score=6, issues=[
                {"section_title": "Risks", "problem": "still weak", "fix_instruction": "add more"}])
    r = DocumentAgent(FakeLLM(), tmp_path, reviewer_llm=StubbornReviewer()).run(COMPLEX_REQUEST)
    assert r.reflection.final_score == 6 and r.reflection.remaining_issues[0].problem == "still weak"
    assert r.status == "completed_with_warnings"


def test_cost_notes_keep_effort_but_drop_money():
    from app.tools import _strip_figures
    assert _strip_figures("Total is $48,480 plus 10% contingency. Lean team over 6 weeks.") == \
        "Lean team over 6 weeks."


def test_cost_estimate_rejects_unknown_non_labour_item():
    with pytest.raises(ValueError, match="unknown item"):
        CostEstimate(line_items=[{"role": "QA Engineer", "days": 5}, {"role": "Trainer", "days": 1}],
                     other_items=[{"item": "Gold-plated servers", "quantity": 1}])


def test_reviewer_outage_returns_unreviewed_document(tmp_path):
    """A provider error during the self-check must not throw away the drafted document."""
    class DownReviewer(FakeLLM):
        def generate(self, system, user, schema):
            raise LLMError("provider rate limit")
    r = DocumentAgent(FakeLLM(), tmp_path, reviewer_llm=DownReviewer()).run(COMPLEX_REQUEST)
    assert r.reflection.skipped and "rate limit" in r.reflection.skipped_reason
    assert r.status == "completed_with_warnings" and "unreviewed" in r.summary
    assert (tmp_path / f"{r.request_id}.docx").exists()


def test_provider_rate_limit_maps_to_429_with_retry_after(make_client):
    import httpx
    from openai import RateLimitError

    class Limited:
        def run(self, request):
            resp = httpx.Response(429, headers={"retry-after": "120"},
                                  request=httpx.Request("POST", "http://x"))
            raise RateLimitError("Rate limit reached: tokens per day", response=resp, body=None)

    app.dependency_overrides[get_agent] = lambda: Limited()
    r = TestClient(app).post("/agent", json={"request": COMPLEX_REQUEST})
    assert r.status_code == 429 and r.headers["retry-after"] == "120"
    assert "tokens per day" in r.json()["detail"]


def test_tool_failure_is_isolated(make_client):
    client = make_client(FakeLLM(fail_section="Timeline", review_fails_first=False))
    body = client.post("/agent", json={"request": COMPLEX_REQUEST}).json()
    timeline = next(t for t in body["plan"] if t["section_title"] == "Timeline")
    assert timeline["status"] == "failed" and "simulated" in timeline["error"]
    assert body["status"] == "completed_with_warnings"


def test_transient_failure_recovered_on_retry(make_client):
    client = make_client(FakeLLM(fail_once="Timeline", review_fails_first=False))
    body = client.post("/agent", json={"request": COMPLEX_REQUEST}).json()
    assert all(t["status"] == "done" for t in body["plan"])
    execute = next(s for s in body["execution_log"] if s["step"] == "execute")
    assert "recovered on retry" in execute["detail"] and "Timeline" in execute["detail"]
    assert body["status"] == "completed"


def test_numeric_table_cells_are_coerced():
    s = Section(title="Budget", blocks=[{"type": "table", "headers": ["Item", 2026],
                                         "rows": [["Licences", 2500], ["Contingency", None]]}])
    assert s.blocks[0].headers == ["Item", "2026"]
    assert s.blocks[0].rows == [["Licences", "2500"], ["Contingency", ""]]


def test_cost_estimate_rejects_roles_not_on_rate_card():
    with pytest.raises(ValueError, match="unknown role"):
        CostEstimate(line_items=[{"role": "Wizard", "days": 5}, {"role": "QA Engineer", "days": 5}])


def test_separate_reviewer_model_does_the_review(tmp_path):
    writer, reviewer = FakeLLM(review_fails_first=False), FakeLLM(review_fails_first=False)
    DocumentAgent(writer, tmp_path, reviewer_llm=reviewer).run(COMPLEX_REQUEST)
    assert "Review" not in [n for n, _ in writer.calls]
    assert [n for n, _ in reviewer.calls] == ["Review"]


def test_guardrail_rejects_non_document_request(make_client):
    r = make_client(FakeLLM(accept=False)).post("/agent", json={"request": "tell me a joke about cats please"})
    assert r.status_code == 422


def test_input_validation(make_client):
    client = make_client(FakeLLM())
    assert client.post("/agent", json={"request": "   hi    "}).status_code == 422
    assert client.post("/agent", json={}).status_code == 422


def test_download_rejects_bad_ids():
    client = TestClient(app)
    assert client.get("/documents/..%2F..%2Fetc").status_code in (400, 404)
    assert client.get("/documents/abcdef1234").status_code == 404
    assert client.get("/runs/not-an-id").status_code == 400


# ------------------------------------------------------------- LLM layer


def test_extract_json_handles_fences_and_chatter():
    assert _extract_json('Sure! ```json\n{"a": 1}\n``` hope that helps') == '{"a": 1}'
    assert _extract_json('Here you go: {"a": {"b": 2}} done') == '{"a": {"b": 2}}'


def _fake_completion(content):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])


def test_llm_retries_with_validation_feedback():
    llm = OpenAICompatibleLLM("http://x", "k", "m", json_retries=2)
    outputs = iter(['{"passed": true}', '{"passed": true, "score": 9, "issues": []}'])
    seen = []

    def create(**kw):
        seen.append(kw["messages"])
        return _fake_completion(next(outputs))

    llm._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    review = llm.generate("sys", "user", Review)
    assert review.score == 9
    assert "invalid" in seen[1][-1]["content"]  # error fed back on the retry


def _bad_request(code, failed_generation=""):
    import httpx
    from openai import BadRequestError
    resp = httpx.Response(400, request=httpx.Request("POST", "http://x"))
    return BadRequestError("bad", response=resp, body={"error": {
        "code": code, "failed_generation": failed_generation}})


def test_llm_recovers_from_provider_json_rejection():
    """Groq returns 400 json_validate_failed instead of content; we must retry, not crash."""
    llm = OpenAICompatibleLLM("http://x", "k", "m", json_retries=2)
    seen = []

    def create(**kw):
        seen.append(kw["messages"])
        if len(seen) == 1:
            raise _bad_request("json_validate_failed", '{"passed": true, "score": 9, "issues": [')
        return _fake_completion('{"passed": true, "score": 9, "issues": []}')

    llm._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    assert llm.generate("sys", "user", Review).score == 9
    assert seen[1][-2]["content"].startswith('{"passed"')  # broken output fed back for correction


def test_llm_reraises_other_bad_requests():
    llm = OpenAICompatibleLLM("http://x", "k", "m")

    def create(**kw):
        raise _bad_request("model_not_found")

    llm._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    from openai import BadRequestError
    with pytest.raises(BadRequestError):
        llm.generate("sys", "user", Review)


def test_provider_errors_are_metered():
    import httpx
    from openai import RateLimitError
    from app.llm import Meter
    llm = OpenAICompatibleLLM("http://x", "k", "m")

    def create(**kw):
        resp = httpx.Response(429, request=httpx.Request("POST", "http://x"))
        raise RateLimitError("slow down", response=resp, body=None)

    llm._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    meter = Meter()
    with pytest.raises(RateLimitError):
        llm.generate("sys", "user", Review, meter=meter)
    assert meter.usage.api_errors == 1 and meter.usage.api_requests == 1


def test_llm_gives_up_after_retries():
    llm = OpenAICompatibleLLM("http://x", "k", "m", json_retries=1)
    llm._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
        create=lambda **kw: _fake_completion("not json at all"))))
    with pytest.raises(LLMError):
        llm.generate("sys", "user", Review)
